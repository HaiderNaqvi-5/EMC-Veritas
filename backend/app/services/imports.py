from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.domain import Activity, ActivityParticipant, Student
from app.services.students import normalize_roll_number

ROLL = {"roll number", "roll no", "registration no", "roll_number"}
NAME = {"full name", "student name", "name", "full_name"}

def preview_students(db: Session, content: bytes) -> dict:
    sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
    headers = {str(value).strip().lower(): index for index, value in enumerate(next(sheet.iter_rows(values_only=True))) if value}
    roll_index = next((headers[key] for key in ROLL if key in headers), None); name_index = next((headers[key] for key in NAME if key in headers), None)
    if roll_index is None or name_index is None: raise ValueError("Spreadsheet must include Roll Number and Full Name headings")
    seen=set(); rows=[]; counts={"valid_rows":0,"duplicate_rows":0,"conflicting_rows":0,"invalid_rows":0}
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        roll=normalize_roll_number(str(values[roll_index])) if values[roll_index] is not None else ""; name=str(values[name_index]).strip() if values[name_index] is not None else ""
        outcome="valid"; detail=None
        if not roll or not name: outcome="invalid"; detail="Roll number and full name are required"
        elif roll in seen: outcome="duplicate"; detail="Duplicate roll number in spreadsheet"
        else:
            existing=db.scalar(select(Student).where(Student.roll_number == roll))
            if existing and existing.full_name != name: outcome="conflict"; detail="Existing student name differs; choose a resolution explicitly"
        seen.add(roll)
        count_key = "conflicting_rows" if outcome == "conflict" else f"{outcome}_rows"
        counts[count_key] += 1
        rows.append({"row_number":number,"roll_number":roll or None,"full_name":name or None,"outcome":outcome,"detail":detail})
    return {**counts,"rows":rows}

def participant_export(db: Session, activity_id) -> bytes:
    activity = db.get(Activity, activity_id)
    if activity is None: raise LookupError("Activity not found")
    rows = db.execute(select(ActivityParticipant, Student).join(Student, ActivityParticipant.student_id == Student.id).where(ActivityParticipant.activity_id == activity_id)).all()
    workbook=Workbook(); sheet=workbook.active; sheet.title="Participants"
    sheet.append(["Roll Number","Student Name","Activity Name","Activity Date","Eligibility Status"])
    for participant, student in rows: sheet.append([student.roll_number,student.full_name,activity.name,activity.activity_date.isoformat(),"Eligible" if participant.eligible else "Ineligible"])
    output=BytesIO(); workbook.save(output); return output.getvalue()

def import_activity_participants(db: Session, activity_id, content: bytes) -> dict:
    """Link existing active students from an Excel attendee list to one activity."""
    activity = db.get(Activity, activity_id)
    if activity is None: raise LookupError("Activity not found")
    sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
    headers = {str(value).strip().lower(): index for index, value in enumerate(next(sheet.iter_rows(values_only=True))) if value}
    roll_index = next((headers[key] for key in ROLL if key in headers), None)
    if roll_index is None: raise ValueError("Spreadsheet must include a Roll Number heading")
    added = existing = 0; unknown_roll_numbers: list[str] = []; seen: set[str] = set()
    for values in sheet.iter_rows(min_row=2, values_only=True):
        roll = normalize_roll_number(str(values[roll_index])) if values[roll_index] is not None else ""
        if not roll or roll in seen: continue
        seen.add(roll)
        student = db.scalar(select(Student).where(Student.roll_number == roll, Student.active.is_(True)))
        if student is None:
            unknown_roll_numbers.append(roll); continue
        participant = db.scalar(select(ActivityParticipant).where(ActivityParticipant.activity_id == activity_id, ActivityParticipant.student_id == student.id))
        if participant is not None:
            existing += 1; continue
        db.add(ActivityParticipant(activity_id=activity_id, student_id=student.id, eligible=True)); added += 1
    return {"added": added, "already_present": existing, "unknown_roll_numbers": unknown_roll_numbers}
